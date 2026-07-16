# -*- coding: utf-8 -*-
"""
ETL: Notes people (+ optional roster Excel) -> SQLite.employees

Purpose:
- Sync active employees from Notes View-LeaveData
- Infer foreign_type / shift_cat from SHIFT_DESC
- Optional: use roster workbook (點將錄-2571/2572) to auto-mark absent people inactive
- Optional: exclude emp_no starts with Q (intern)

Usage:
  python backend/etl/etl_notes_people.py
  python backend/etl/etl_notes_people.py --dept 2531 2537 --mark-inactive
  python backend/etl/etl_notes_people.py --schedule "D:\\Program\\Notes\\2026_0620_0820_2571-2572出勤表.xlsx" --mark-inactive
"""

from __future__ import annotations

import argparse
import os
import re
import sqlite3
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple

import win32com.client
from openpyxl import load_workbook


DEFAULT_DEPTS = ["2531", "2537"]
DEFAULT_SCHEDULE = "2026_0620_0820_2571-2572出勤表.xlsx"
CFG_FILENAME = "部門設定_請先改這裡.ps1"


def get_db_path(cli_db: str) -> Path:
    if cli_db:
        p = Path(cli_db)
        if p.is_absolute():
            return p
        return Path(__file__).resolve().parents[2] / p

    env = os.getenv("MOS_DB_PATH", "").strip()
    if env:
        return Path(env)

    return Path(r"D:\Fastapi\src\data\mos_attendance.db")


def parse_settings() -> Tuple[List[str], str]:
    depts = DEFAULT_DEPTS[:]
    notes_password = ""

    cur = Path(__file__).resolve()
    candidates = [
        cur.parent / CFG_FILENAME,
        cur.parents[1] / CFG_FILENAME,
        cur.parents[2] / CFG_FILENAME,
    ]

    found = None
    for p in candidates:
        if p.exists():
            found = p
            break

    if not found:
        return depts, notes_password

    text = found.read_text(encoding="utf-8", errors="ignore")
    m = re.search(r"\$Depts\s*=\s*@\((.*?)\)", text, re.S)
    if m:
        vals = re.findall(r'"([^"]+)"', m.group(1))
        if vals:
            depts = vals

    m2 = re.search(r'\$NOTES_PASSWORD\s*=\s*"([^"]*)"', text)
    if m2:
        notes_password = m2.group(1)

    return depts, notes_password


def get_item(doc, name: str):
    try:
        vals = doc.GetItemValue(name)
        if vals and len(vals) > 0:
            return vals[0]
    except Exception:
        pass
    return None


def infer_foreign_type(shift_desc: str) -> str:
    return "外勞" if "外勞" in (shift_desc or "") else "本勞"


def infer_shift_cat(shift_desc: str) -> str:
    return "常日" if "常日" in (shift_desc or "") else "輪班"


def fetch_notes_people(depts: Iterable[str], notes_password: str, exclude_q: bool) -> Dict[str, Dict[str, str]]:
    print("[1/3] Fetch people from Notes ...")
    sess = win32com.client.Dispatch("Notes.NotesSession")
    if hasattr(sess, "Initialize"):
        try:
            sess.Initialize(notes_password or "")
        except Exception:
            pass

    db = sess.GetDatabase("CMOS07/CHIPMOS", "ap\\office\\pds0001.nsf")
    if not db.IsOpen:
        db.Open("CMOS07/CHIPMOS", "ap\\office\\pds0001.nsf")

    dt1900 = sess.CreateDateTime("01/01/1900 00:00:00")
    formula = 'SELECT @All & Form != ""'
    dc = db.Search(formula, dt1900, 0)
    if dc is None:
        return {}

    people: Dict[str, Dict[str, str]] = {}
    depts_set = set(str(x) for x in depts)

    total = dc.Count
    for i in range(1, total + 1):
        doc = dc.GetNthDocument(i)
        if doc is None:
            continue

        dept_no = str(get_item(doc, "DeptNo") or "").strip()
        if dept_no not in depts_set:
            continue

        emp_no = str(get_item(doc, "EmpNo") or "").strip()
        if not emp_no:
            continue
        if exclude_q and emp_no.startswith("Q"):
            continue

        emp_name = str(get_item(doc, "EmpName") or "").strip()
        shift_desc = str(get_item(doc, "SHIFT_DESC") or "").strip()

        people[emp_no] = {
            "emp_no": emp_no,
            "emp_name": emp_name,
            "dept_no": dept_no,
            "shift_raw": shift_desc,
            "foreign_type": infer_foreign_type(shift_desc),
            "shift_cat": infer_shift_cat(shift_desc),
        }

        if i % 1000 == 0:
            print(f"    scanned {i}/{total}, collected {len(people)}")

    print(f"    collected from Notes: {len(people)}")
    return people


def parse_roster_empnos(schedule_path: Optional[Path]) -> Set[str]:
    if not schedule_path:
        return set()
    if not schedule_path.exists():
        print(f"    ⚠ roster file not found: {schedule_path} (skip roster filter)")
        return set()

    print("[2/3] Parse roster workbook ...")
    wb = load_workbook(schedule_path, data_only=True)
    empnos: Set[str] = set()
    for sheet_name in ("點將錄-2571", "點將錄-2572"):
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for r in range(3, ws.max_row + 1):
            emp_no = str(ws.cell(r, 4).value or "").strip()
            shift = str(ws.cell(r, 9).value or "").strip()
            if not emp_no:
                continue
            if not re.match(r"^[QS]\d+$", emp_no):
                continue
            if not shift:
                continue
            empnos.add(emp_no)
    wb.close()
    print(f"    roster emp count: {len(empnos)}")
    return empnos


def upsert_employee(conn: sqlite3.Connection, p: Dict[str, str]):
    conn.execute(
        """
        INSERT INTO employees(emp_no, emp_name, dept_no, foreign_type, shift_cat, shift_raw, active, updated_at)
        VALUES(?, ?, ?, ?, ?, ?, 1, datetime('now'))
        ON CONFLICT(emp_no) DO UPDATE SET
            emp_name=excluded.emp_name,
            dept_no=excluded.dept_no,
            foreign_type=excluded.foreign_type,
            shift_cat=excluded.shift_cat,
            shift_raw=excluded.shift_raw,
            active=1,
            updated_at=datetime('now')
        """,
        (
            p["emp_no"], p["emp_name"], p["dept_no"],
            p["foreign_type"], p["shift_cat"], p["shift_raw"]
        )
    )


def mark_inactive_not_in_latest(conn: sqlite3.Connection, latest_empnos: Set[str]):
    rows = conn.execute("SELECT emp_no FROM employees WHERE active=1").fetchall()
    current_active = {r[0] for r in rows}

    to_disable = current_active - latest_empnos
    for emp_no in to_disable:
        conn.execute(
            "UPDATE employees SET active=0, updated_at=datetime('now') WHERE emp_no=?",
            (emp_no,)
        )
    return len(to_disable)


def sync_people(
    db_path: Path,
    notes_people: Dict[str, Dict[str, str]],
    roster_empnos: Set[str],
    use_roster_filter: bool,
    mark_inactive: bool,
):
    print("[3/3] Sync into SQLite.employees ...")
    db_path.parent.mkdir(parents=True, exist_ok=True)

    latest: Dict[str, Dict[str, str]] = {}
    if use_roster_filter and roster_empnos:
        for emp_no, p in notes_people.items():
            if emp_no in roster_empnos:
                latest[emp_no] = p
    else:
        latest = notes_people

    conn = sqlite3.connect(db_path)
    try:
        conn.execute("PRAGMA foreign_keys=ON;")
        conn.execute("PRAGMA journal_mode=WAL;")
        conn.execute("PRAGMA synchronous=NORMAL;")

        for p in latest.values():
            upsert_employee(conn, p)

        disabled = 0
        if mark_inactive:
            disabled = mark_inactive_not_in_latest(conn, set(latest.keys()))

        conn.commit()
        print(f"    upserted: {len(latest)}")
        if mark_inactive:
            print(f"    set inactive: {disabled}")
    finally:
        conn.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--db", default="", help="SQLite path override")
    ap.add_argument("--dept", nargs="*", default=None, help="Dept override")
    ap.add_argument("--schedule", default=DEFAULT_SCHEDULE, help="Roster workbook path")
    ap.add_argument("--no-roster-filter", action="store_true", help="Do not restrict by roster emp list")
    ap.add_argument("--mark-inactive", action="store_true", help="Mark employees inactive if not in latest result")
    ap.add_argument("--include-q", action="store_true", help="Include emp_no starting with Q")
    args = ap.parse_args()

    cfg_depts, notes_password = parse_settings()
    depts = args.dept if args.dept else cfg_depts
    db_path = get_db_path(args.db)

    sch = Path(args.schedule)
    if not sch.is_absolute():
        sch = Path(__file__).resolve().parents[2] / sch

    print("===================================================")
    print(" ETL Notes People -> SQLite")
    print("===================================================")
    print(f"DB            : {db_path}")
    print(f"Depts         : {', '.join(depts)}")
    print(f"Schedule      : {sch}")
    print(f"Use roster    : {not args.no_roster_filter}")
    print(f"Mark inactive : {args.mark_inactive}")
    print(f"Include Q     : {args.include_q}")

    notes_people = fetch_notes_people(depts, notes_password, exclude_q=(not args.include_q))
    roster_empnos = parse_roster_empnos(sch)
    sync_people(
        db_path=db_path,
        notes_people=notes_people,
        roster_empnos=roster_empnos,
        use_roster_filter=(not args.no_roster_filter),
        mark_inactive=args.mark_inactive,
    )
    print("Done.")


if __name__ == "__main__":
    main()
